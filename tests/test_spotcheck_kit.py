import pytest
from openpyxl import load_workbook

from src.evaluation import analyze_corrected_rerun_spotcheck as analysis
from src.evaluation import build_spotcheck_evaluator_kit as kit
from src.evaluation import import_spotcheck_workbook as importer
from src.evaluation.build_corrected_rerun_spotcheck import select_questions


@pytest.fixture()
def built_kit(tmp_path, monkeypatch):
    monkeypatch.setattr(kit, "OUT", tmp_path / "kit")
    monkeypatch.setattr(kit, "WORKBOOK", tmp_path / "kit" / "scoring.xlsx")
    monkeypatch.setattr(kit, "PDF", tmp_path / "kit" / "guidelines.pdf")
    monkeypatch.setattr(kit, "CSV_COPY", tmp_path / "kit" / "blank.csv")
    monkeypatch.setattr(kit, "ZIP", tmp_path / "kit.zip")
    monkeypatch.setattr(importer, "FILLED", tmp_path / "filled.csv")
    kit.main()
    return tmp_path


def _fill(path, value=3, flag=0):
    wb = load_workbook(path)
    ws = wb["Scoring"]
    for r in range(2, 62):
        for c in range(8, 12):
            ws.cell(r, c, value)
        ws.cell(r, 12, flag)
    return wb, ws


def test_filled_workbook_round_trips_into_analysis(built_kit):
    wb, _ = _fill(built_kit / "kit" / "scoring.xlsx")
    wb.save(built_kit / "done.xlsx")
    importer.main(built_kit / "done.xlsx", force=False)
    scores = analysis.read_scores(built_kit / "filled.csv", select_questions())
    assert len(scores) == 60 and all(s["correctness"] == 3 for s in scores.values())


def test_invalid_score_is_rejected(built_kit):
    wb, ws = _fill(built_kit / "kit" / "scoring.xlsx")
    ws.cell(5, 8, 7)
    wb.save(built_kit / "bad.xlsx")
    with pytest.raises(ValueError, match="correctness"):
        importer.main(built_kit / "bad.xlsx", force=False)


def test_edited_answer_text_is_rejected(built_kit):
    wb, ws = _fill(built_kit / "kit" / "scoring.xlsx")
    ws.cell(2, 7, "edited answer")
    wb.save(built_kit / "edited.xlsx")
    with pytest.raises(ValueError, match="text differs"):
        importer.main(built_kit / "edited.xlsx", force=False)
