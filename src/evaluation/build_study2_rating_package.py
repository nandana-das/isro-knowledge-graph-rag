"""Build the Study 2 rating package and evaluator kit (preregistration_study2_generators.md §8).

60 KG-required questions x 6 candidates (V and A x three generators), shuffled
per question and labelled A-F. Writes:
- data/annotations/study2_rating_blank.csv: the blank package (identical for every rater)
- data/results/study2_generators/rating_mapping.json: label -> (generator, system); raters never see it
- handoff/study2_rater/: Scoring workbook, guidelines PDF and CSV backup, zipped

    python -m src.evaluation.build_study2_rating_package
"""

from __future__ import annotations

import csv
import json
import random
import shutil
import zipfile
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font

from src.evaluation import build_spotcheck_evaluator_kit as kit
from src.evaluation.build_corrected_rerun_human_package import (
    ANNOTATIONS, BENCHMARK, CHUNKS, SEED, TRIPLES, answer_hash, evidence_pack, load_jsonl, sha256_bytes,
)
from src.evaluation.run_generator_study import GENERATORS, OUTPUT, ROOT, SYSTEMS, results_path

LABELS = ("A", "B", "C", "D", "E", "F")
DIMENSIONS = ("correctness", "completeness", "groundedness", "relevance", "unsupported_claim")
BLANK = ANNOTATIONS / "study2_rating_blank.csv"
MAPPING = OUTPUT / "rating_mapping.json"
PACKAGE_MANIFEST = OUTPUT / "rating_package_manifest.json"
KIT = ROOT / "handoff" / "study2_rater"
WORKBOOK = KIT / "KG-RAG_Study2_Scoring.xlsx"
PDF = KIT / "KG-RAG_Study2_Evaluator_Guidelines.pdf"
ZIP = ROOT / "handoff" / "KG-RAG_Study2_Evaluator_Package.zip"

# Rubric wording is the Study 1 rubric (kit.SECTIONS "Scores" and "Unsupported claim"), unchanged.
_RUBRIC = {name: paragraphs for name, paragraphs in kit.SECTIONS}
SECTIONS = [
    ("What this is", [
        "You are evaluating answers produced by question-answering systems about Indian space missions "
        "(ISRO/ISSDC). There are 60 questions. Each has six candidate answers, labelled A to F, "
        "so you will score 360 answers in total. Expect roughly 6 to 8 hours; you may split the work into several sessions.",
        "Two other people score the same answers independently. It matters that you work completely on your own.",
    ]),
    ("Before you start (independence rules)", [
        "- Do not look at any other ratings, scores, results or analysis files for this project.",
        "- Do not use ChatGPT, Claude or any other AI tool to read, score or suggest scores.",
        "- Do not discuss the questions or answers with anyone until all three evaluators have finished.",
        "- Candidate letters are shuffled for every question. Do not try to work out which system wrote which answer.",
        "- You do not need the internet. Everything you need is in the workbook.",
    ]),
    ("What you get for each question", [
        "- The question.",
        "- The reference answer and acceptable answers.",
        "- An evidence pack: supporting facts and the official ISRO/ISSDC source text. It is the same for all six answers.",
        "- Six candidate answers (A to F). Some may be identical; score each one.",
    ]),
    ("Scores (1-5)", _RUBRIC["Scores (1-5)"]),
    ("Unsupported claim (0/1)", _RUBRIC["Unsupported claim (0/1)"]),
    ("How to fill in the workbook", [
        "1. Open the Scoring sheet. Each row is one answer; the question, reference and evidence pack span its six rows.",
        "2. For each answer, choose a value in the five yellow columns (Correctness, Completeness, Groundedness, "
        "Relevance, Unsupported claim) from the dropdown. Blank score cells are highlighted red until you fill them.",
        "3. Fill every score cell. If an answer genuinely cannot be judged, still score it and write the reason in the Notes column.",
        "4. Do not edit the question, reference, evidence or answer text. Those cells are locked to prevent accidents.",
        "5. For every working session, write the date, start time and end time on the Session log sheet.",
        "6. Save the workbook in .xlsx format and send it back. Please add your initials to the file name, "
        "e.g. KG-RAG_Study2_Scoring_AB.xlsx.",
    ]),
]


def build_package() -> list[dict]:
    questions = [q for q in json.loads(BENCHMARK.read_text(encoding="utf8"))["questions"] if q["kg_required"] == "YES"]
    chunks = {c["chunk_id"]: c for c in load_jsonl(CHUNKS)}
    triples = {t["triple_id"]: t for t in load_jsonl(TRIPLES)}
    answers = {}
    for generator in GENERATORS:
        for row in load_jsonl(results_path(generator)):
            answers[(row["question_id"], generator, row["system"])] = row["answer"]
    candidates = [(g, s) for g in GENERATORS for s in SYSTEMS]
    missing = [(q["question_id"], g, s) for q in questions for g, s in candidates if (q["question_id"], g, s) not in answers]
    if missing:
        raise RuntimeError(f"Generations missing: {missing[:5]} ({len(missing)} total)")

    rng = random.Random(SEED)
    rows, mapping = [], {}
    for q in questions:
        qid = q["question_id"]
        order = list(candidates)
        rng.shuffle(order)
        label_to_candidate = dict(zip(LABELS, order))
        mapping[qid] = {
            "label_to_candidate": {label: {"generator": g, "system": s} for label, (g, s) in label_to_candidate.items()},
            "label_hashes": {label: answer_hash(answers[(qid, g, s)]) for label, (g, s) in label_to_candidate.items()},
        }
        row = {
            "question_id": qid,
            "question": q["question"],
            "reference_answer": q["reference_answer"],
            "acceptable_answers": " | ".join(q["acceptable_answers"]),
            "evidence_pack": evidence_pack(q, chunks, triples),
        }
        for label in LABELS:
            g, s = label_to_candidate[label]
            row[f"candidate_{label}"] = answers[(qid, g, s)]
        for dimension in DIMENSIONS:
            for label in LABELS:
                row[f"{label}_{dimension}"] = ""
        rows.append(row)

    ANNOTATIONS.mkdir(parents=True, exist_ok=True)
    with BLANK.open("w", encoding="utf8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    MAPPING.write_text(json.dumps({"seed": SEED, "labels": list(LABELS), "mapping": mapping}, indent=2) + "\n", encoding="utf8")
    PACKAGE_MANIFEST.write_text(json.dumps({
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "questions": len(rows),
        "candidates": len(rows) * len(LABELS),
        "blank_sha256": sha256_bytes(BLANK.read_bytes()),
        "mapping_sha256": sha256_bytes(MAPPING.read_bytes()),
        "results_sha256": {g: sha256_bytes(results_path(g).read_bytes()) for g in GENERATORS},
        "warning": "Raters must not open the mapping file.",
    }, indent=2) + "\n", encoding="utf8")
    return rows


def build_kit(rows: list[dict]) -> None:
    KIT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    kit.build_instructions(wb.active, SECTIONS)
    kit.build_scoring(wb.create_sheet(), rows, LABELS)
    log = wb.create_sheet("Session log")
    for col, (title, width) in enumerate((("Date", 14), ("Start time", 12), ("End time", 12), ("Questions covered", 22), ("Notes", 40)), 1):
        log.cell(1, col, title).font = Font(name=kit.ARIAL, size=10, bold=True)
        log.column_dimensions[log.cell(1, col).column_letter].width = width
    wb.save(WORKBOOK)
    kit.build_pdf(PDF, SECTIONS)
    shutil.copyfile(BLANK, KIT / BLANK.name)
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in (PDF, WORKBOOK, KIT / BLANK.name):
            archive.write(path, path.name)


def main() -> None:
    manifest = json.loads((OUTPUT / "run_manifest.json").read_text(encoding="utf8"))
    print(f"generation manifest: {manifest['code_commit']}")
    rows = build_package()
    build_kit(rows)
    print(f"built {len(rows)} questions x {len(LABELS)} candidates; kit at {ZIP}")


if __name__ == "__main__":
    main()
