"""Convert the evaluator's filled spot-check workbook into the analysis CSV.

Reads the Scoring sheet of KG-RAG_Spotcheck_Scoring*.xlsx, checks every answer
text against the blank package and every score against the allowed values,
and writes data/annotations/corrected_rerun_spotcheck_filled.csv.

    python -m src.evaluation.import_spotcheck_workbook path/to/filled.xlsx
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

from src.evaluation.build_corrected_rerun_spotcheck import BLANK, FILLED

LABELS = ("A", "B", "C")
SCORES = ("correctness", "completeness", "groundedness", "relevance", "unsupported_claim")
FIRST_SCORE_COLUMN = 8  # H


def norm(text: str) -> str:
    return " ".join((text or "").split())


def read_workbook(path: Path) -> dict[tuple[str, str], dict]:
    ws = load_workbook(path, data_only=True)["Scoring"]
    entries: dict[tuple[str, str], dict] = {}
    question_id = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1]:
            question_id = str(row[1])
        label = row[5]
        if label not in LABELS:
            continue
        scores = {}
        for offset, name in enumerate(SCORES):
            raw = row[FIRST_SCORE_COLUMN - 1 + offset]
            value = int(raw) if isinstance(raw, (int, float)) and float(raw).is_integer() else (
                int(raw) if isinstance(raw, str) and raw.strip().isdigit() else None)
            allowed = (0, 1) if name == "unsupported_claim" else (1, 2, 3, 4, 5)
            if value not in allowed:
                raise ValueError(f"{question_id} answer {label}: {name} is {raw!r}; expected one of {allowed}")
            scores[name] = value
        entries[(question_id, label)] = {"answer": row[6] or "", "scores": scores, "notes": row[12] or ""}
    return entries


def main(workbook: Path, force: bool) -> None:
    if FILLED.exists() and not force:
        raise FileExistsError(f"{FILLED} already exists; pass --force to replace it")
    with BLANK.open(encoding="utf8", newline="") as handle:
        reader = csv.DictReader(handle)
        fields, rows = reader.fieldnames, list(reader)
    entries = read_workbook(workbook)
    expected = {(r["question_id"], label) for r in rows for label in LABELS}
    if set(entries) != expected:
        raise ValueError(f"Workbook answers do not match the package: missing {sorted(expected - set(entries))}")
    notes = []
    for row in rows:
        for label in LABELS:
            entry = entries[(row["question_id"], label)]
            if norm(entry["answer"]) != norm(row[f"candidate_{label}"]):
                raise ValueError(f"{row['question_id']} answer {label}: text differs from the package")
            for name, value in entry["scores"].items():
                row[f"{label}_{name}"] = value
            if entry["notes"]:
                notes.append(f"{row['question_id']} {label}: {entry['notes']}")
    with FILLED.open("w", encoding="utf8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    print(f"wrote {FILLED.name}: {len(rows)} questions, {len(entries) * len(SCORES)} scores")
    if notes:
        print("Evaluator notes:\n  " + "\n  ".join(notes))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    main(args.workbook, args.force)
