"""Build the hand-off kit for the independent spot-check evaluator (preregistration §14).

Creates handoff/spotcheck_evaluator/ with:
- KG-RAG_Spotcheck_Scoring.xlsx: Instructions sheet + Scoring sheet (one row per answer)
- KG-RAG_Spotcheck_Evaluator_Guidelines.pdf: the same guidelines, printable
- corrected_rerun_spotcheck.csv: plain-text backup of the blank package
and zips them. The scoring rubric is the one rater 1 received
(data/annotations/corrected_rerun_human_eval_guidelines.md), unchanged.

    python -m src.evaluation.build_spotcheck_evaluator_kit
"""

from __future__ import annotations

import csv
import shutil
import zipfile
from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

from src.evaluation.build_corrected_rerun_spotcheck import BLANK

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "handoff" / "spotcheck_evaluator"
WORKBOOK = OUT / "KG-RAG_Spotcheck_Scoring.xlsx"
PDF = OUT / "KG-RAG_Spotcheck_Evaluator_Guidelines.pdf"
CSV_COPY = OUT / "corrected_rerun_spotcheck.csv"
ZIP = ROOT / "handoff" / "KG-RAG_Spotcheck_Evaluator_Package.zip"
LABELS = ("A", "B", "C")
SCORES = ("correctness", "completeness", "groundedness", "relevance", "unsupported_claim")

# ---- Guidelines: one source for the PDF and the workbook ----
# Rubric wording follows corrected_rerun_human_eval_guidelines.md exactly.
TITLE = "KG-RAG Answer Evaluation: Evaluator Guidelines"
SECTIONS: list[tuple[str, list[str]]] = [
    ("What this is", [
        "You are evaluating answers produced by question-answering systems about Indian space missions "
        "(ISRO/ISSDC). There are 20 questions. Each has three candidate answers, labelled A, B and C, "
        "so you will score 60 answers in total. It takes roughly 60 to 90 minutes.",
        "Your scores are used to check another set of ratings, so it matters that you work completely on your own.",
    ]),
    ("Before you start (independence rules)", [
        "- Do not look at any other ratings, scores, results or analysis files for this project.",
        "- Do not use ChatGPT, Claude or any other AI tool to read, score or suggest scores.",
        "- Do not discuss the questions or answers with anyone until you have finished.",
        "- Candidate letters are shuffled for every question. Do not try to work out which system wrote which answer.",
        "- You do not need the internet. Everything you need is in the workbook.",
    ]),
    ("What you get for each question", [
        "- The question.",
        "- The reference answer and acceptable answers.",
        "- An evidence pack: supporting facts and the official ISRO/ISSDC source text. It is the same for all three answers.",
        "- Three candidate answers (A, B, C).",
    ]),
    ("Scores (1-5)", [
        "Correctness: 1 substantially incorrect, 2 mostly incorrect, 3 partially correct, 4 mostly correct, 5 fully correct. "
        "Some questions have several correct answers (e.g. \"Which organization developed a payload carried by Aditya-L1?\"). "
        "The reference is one example: any answer the evidence pack supports as a member of the correct set counts as correct. "
        "Do not penalise wording, acronym, alias or date-format differences.",
        "Completeness: 1 misses essentially all required information, 3 partial, 5 complete.",
        "Groundedness: judged against the evidence pack. 1 largely unsupported, 3 partially supported, 5 fully supported. "
        "A correct \"I don't know\" makes no claim and is fully grounded, but scores low on correctness and completeness.",
        "Relevance: 1 largely irrelevant, 3 partially relevant, 5 directly relevant.",
        "Use 2 and 4 for answers that fall between the described levels.",
    ]),
    ("Unsupported claim (0/1)", [
        "1 if the answer makes at least one material claim that the evidence pack does not support, otherwise 0.",
    ]),
    ("How to fill in the workbook", [
        "1. Open the Scoring sheet. Each row is one answer; the question, reference and evidence pack span its three rows.",
        "2. For each answer, choose a value in the five yellow columns (Correctness, Completeness, Groundedness, "
        "Relevance, Unsupported claim) from the dropdown. Blank score cells are highlighted red until you fill them.",
        "3. Fill every score cell. If an answer genuinely cannot be judged, still score it and write the reason in the Notes column.",
        "4. Do not edit the question, reference, evidence or answer text. Those cells are locked to prevent accidents.",
        "5. Save the workbook in .xlsx format and send it back. Please add your initials to the file name, "
        "e.g. KG-RAG_Spotcheck_Scoring_AB.xlsx.",
    ]),
]
EXAMPLE = {
    "note": "Format example only. This question is made up and is not part of the evaluation.",
    "question": "Which launch vehicle launched the Example-1 mission?",
    "answer": "Example-1 was launched by the PSLV-C99 launch vehicle.",
    "scores": (5, 5, 5, 5, 0),
}

ARIAL = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="1F3864")
BAND = (PatternFill("solid", fgColor="F2F2F2"), PatternFill("solid", fgColor="FFFFFF"))
RED = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_blank() -> list[dict]:
    with BLANK.open(encoding="utf8", newline="") as handle:
        return list(csv.DictReader(handle))


def lines_needed(text: str, width_chars: float) -> int:
    per_line = max(1, int(width_chars * 1.15))
    return sum(max(1, ceil(len(part) / per_line)) for part in (text or "").split("\n"))


def build_instructions(ws) -> None:
    ws.title = "Instructions"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 120
    row = 1
    ws.cell(row, 1, TITLE).font = Font(name=ARIAL, size=16, bold=True, color="1F3864")
    row += 2
    for heading, paragraphs in SECTIONS:
        ws.cell(row, 1, heading).font = Font(name=ARIAL, size=12, bold=True, color="1F3864")
        row += 1
        for text in paragraphs:
            cell = ws.cell(row, 1, text)
            cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 14 * lines_needed(text, 120)
            row += 1
        row += 1

    ws.cell(row, 1, "Legend").font = Font(name=ARIAL, size=12, bold=True, color="1F3864")
    row += 1
    cell = ws.cell(row, 1, "Yellow cells on the Scoring sheet are the only cells you fill in. All other cells are locked.")
    cell.fill, cell.font = YELLOW, Font(name=ARIAL, size=10)
    row += 1
    cell = ws.cell(row, 1, "Red highlighting means a score cell is still empty.")
    cell.fill, cell.font = RED, Font(name=ARIAL, size=10)
    row += 2

    ws.cell(row, 1, "Example row (format only)").font = Font(name=ARIAL, size=12, bold=True, color="1F3864")
    row += 1
    ws.cell(row, 1, EXAMPLE["note"]).font = Font(name=ARIAL, size=10, italic=True)
    row += 1
    example = (
        f"Question: {EXAMPLE['question']}   |   Answer: {EXAMPLE['answer']}   |   "
        + "   ".join(f"{name.replace('_', ' ').title()}: {value}" for name, value in zip(SCORES, EXAMPLE["scores"]))
    )
    cell = ws.cell(row, 1, example)
    cell.font = Font(name=ARIAL, size=10)
    cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[row].height = 14 * lines_needed(example, 120)


def build_scoring(ws, blank: list[dict]) -> None:
    ws.title = "Scoring"
    headers = ["#", "Question ID", "Question", "Reference / acceptable answers", "Evidence pack",
               "Answer", "Candidate answer", "Correctness (1-5)", "Completeness (1-5)",
               "Groundedness (1-5)", "Relevance (1-5)", "Unsupported claim (0/1)", "Notes (optional)"]
    widths = [5, 11, 28, 28, 70, 8, 55, 13, 13, 13, 13, 14, 28]
    for col, (title, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, col, title)
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = BOX
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    scale = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True, showErrorMessage=True,
                           errorTitle="Invalid score", error="Choose a whole number from 1 to 5.")
    flag = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True,
                          errorTitle="Invalid value", error="Choose 0 (no unsupported claim) or 1 (at least one).")
    ws.add_data_validation(scale)
    ws.add_data_validation(flag)

    row = 2
    for number, item in enumerate(blank, 1):
        band = BAND[number % 2]
        reference = f"Reference: {item['reference_answer']}\n\nAcceptable: {item['acceptable_answers']}"
        shared = {1: number, 2: item["question_id"], 3: item["question"], 4: reference, 5: item["evidence_pack"]}
        for col, value in shared.items():
            ws.merge_cells(start_row=row, start_column=col, end_row=row + 2, end_column=col)
            cell = ws.cell(row, col, value)
            cell.font = Font(name=ARIAL, size=9 if col == 5 else 10, bold=col == 3)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        evidence_lines = max(lines_needed(item["evidence_pack"], widths[4]), lines_needed(reference, widths[3]))
        for offset, label in enumerate(LABELS):
            r = row + offset
            answer = item[f"candidate_{label}"]
            ws.cell(r, 6, label).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(r, 6).font = Font(name=ARIAL, size=11, bold=True)
            cell = ws.cell(r, 7, answer)
            cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col in range(8, 14):
                cell = ws.cell(r, col)
                cell.protection = Protection(locked=False)
                cell.font = Font(name=ARIAL, size=11 if col < 13 else 10)
                cell.alignment = Alignment(horizontal="center" if col < 13 else "left", vertical="top", wrap_text=True)
                if col < 13:
                    cell.fill = YELLOW
            for col in range(1, 14):
                ws.cell(r, col).border = BOX
                if col < 8:
                    ws.cell(r, col).fill = band
            needed = max(lines_needed(answer, widths[6]), ceil(evidence_lines / 3), 3)
            ws.row_dimensions[r].height = min(409, 13 * needed)
        row += 3
    last = row - 1
    scale.add(f"H2:K{last}")
    flag.add(f"L2:L{last}")
    ws.conditional_formatting.add(f"H2:L{last}", FormulaRule(formula=["ISBLANK(H2)"], fill=RED))
    # Lock the sheet without a password so source text cannot be edited by accident.
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def build_pdf() -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=14, spaceAfter=4)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=12, bulletIndent=2)
    heading = ParagraphStyle("h", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12,
                             textColor=colors.HexColor("#1F3864"), spaceBefore=10, spaceAfter=4)
    title = ParagraphStyle("t", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=17,
                           textColor=colors.HexColor("#1F3864"), alignment=0)

    def esc(text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [Paragraph(TITLE, title), Spacer(1, 6)]
    for name, paragraphs in SECTIONS:
        story.append(Paragraph(name, heading))
        for text in paragraphs:
            if text.startswith("- "):
                story.append(Paragraph(esc(text[2:]), bullet, bulletText="•"))
            elif ": " in text and name == "Scores (1-5)" and not text.startswith("Use "):
                label, rest = text.split(": ", 1)
                story.append(Paragraph(f"<b>{esc(label)}:</b> {esc(rest)}", body))
            else:
                story.append(Paragraph(esc(text), body))

    story.append(Paragraph("Quick reference", heading))
    table = Table([
        ["Score", "Correctness", "Completeness", "Groundedness", "Relevance"],
        ["1", "substantially incorrect", "misses essentially all", "largely unsupported", "largely irrelevant"],
        ["2", "mostly incorrect", "between 1 and 3", "between 1 and 3", "between 1 and 3"],
        ["3", "partially correct", "partial", "partially supported", "partially relevant"],
        ["4", "mostly correct", "between 3 and 5", "between 3 and 5", "between 3 and 5"],
        ["5", "fully correct", "complete", "fully supported", "directly relevant"],
    ], colWidths=[1.4 * cm, 3.9 * cm, 3.9 * cm, 3.9 * cm, 3.9 * cm])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
    ]))
    story += [table, Spacer(1, 6), Paragraph("Unsupported claim: 0 = none, 1 = at least one material claim the evidence pack does not support.", body)]
    SimpleDocTemplate(str(PDF), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                      bottomMargin=1.8 * cm, title=TITLE).build(story)


def main() -> None:
    blank = load_blank()
    if len(blank) != 20 or any(v for row in blank for k, v in row.items() if k[:2] in ("A_", "B_", "C_")):
        raise RuntimeError("Expected the blank 20-question spot-check package")
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    build_instructions(wb.active)
    build_scoring(wb.create_sheet(), blank)
    wb.save(WORKBOOK)
    build_pdf()
    shutil.copyfile(BLANK, CSV_COPY)
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in (PDF, WORKBOOK, CSV_COPY):
            archive.write(path, path.name)
    print(f"wrote {OUT} and {ZIP.name}")


if __name__ == "__main__":
    main()
